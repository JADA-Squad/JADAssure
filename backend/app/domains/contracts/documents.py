"""Contract document upload, storage and text extraction (BE-027, FR-13, GR-07).

Files are untrusted. The type is checked by extension AND content signature, size is capped, and
the original is stored through the DocumentStore interface with its sha256. Text is extracted from
PDF (text layer only; scanned images aren't supported in the pilot), DOCX or UTF-8 TXT. The text
is only ever treated as data.
"""

import hashlib
import io
import zipfile

from fastapi import HTTPException, status
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.storage import get_store
from app.core.tenancy import TenantContext
from app.domains.audit import service as audit
from app.domains.contracts.models import Agreement, Document

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 400_000
MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
}


class DocumentError(ValueError):
    """The file can't be accepted; the message is safe to show."""


def _extension(filename: str) -> str:
    dot = filename.rfind(".")
    ext = filename[dot:].lower() if dot >= 0 else ""
    if ext not in MEDIA_TYPES:
        raise DocumentError("Upload a PDF, DOCX or TXT file")
    return ext


def _check_signature(ext: str, data: bytes) -> None:
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise DocumentError("This file isn't a valid PDF")
    if ext == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise DocumentError("This file isn't a valid Word document")
        except zipfile.BadZipFile:
            raise DocumentError("This file isn't a valid Word document") from None
    if ext == ".txt" and b"\x00" in data:
        raise DocumentError("This text file looks binary")


def extract_text(ext: str, data: bytes) -> str:
    if ext == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise DocumentError("Text files must be UTF-8") from None
    elif ext == ".pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except (PdfReadError, ValueError, KeyError):
            raise DocumentError("The PDF couldn't be read") from None
    else:
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip():
        raise DocumentError("No text found. Scanned (image-only) documents aren't supported yet")
    return text[:MAX_TEXT_CHARS]


def upload(session: Session, ctx: TenantContext, agreement: Agreement, filename: str, data: bytes) -> Document:
    try:
        if len(data) > MAX_BYTES:
            raise DocumentError("File is larger than 10 MB")
        ext = _extension(filename)
        _check_signature(ext, data)
        text = extract_text(ext, data)
    except DocumentError as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from None

    digest = hashlib.sha256(data).hexdigest()
    if session.scalar(select(Document.id).where(Document.org_id == ctx.org_id, Document.agreement_id == agreement.id,
                                                Document.sha256 == digest)):
        raise HTTPException(status.HTTP_409_CONFLICT, "This document has already been uploaded for this agreement")
    doc = Document(
        org_id=ctx.org_id, agreement_id=agreement.id, filename=filename[:255], media_type=MEDIA_TYPES[ext],
        storage_key=get_store().put(ctx.org_id, "contracts", data, ext), sha256=digest, size_bytes=len(data),
        uploaded_by=ctx.user_id, text_extracted=text, extraction_status="queued", extraction_stats={},
    )
    session.add(doc)
    session.flush()
    audit.record(session, org_id=ctx.org_id, actor_id=ctx.user_id, action="document.uploaded", entity="document",
                 entity_id=doc.id, payload={"agreement_id": str(agreement.id), "sha256": digest,
                                            "size_bytes": len(data), "media_type": doc.media_type})
    return doc
